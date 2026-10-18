#!/usr/bin/env python3
"""Extrae el texto de todos los archivos de LIBROS-REFERENCIA3 sin hacer checkout.

Lee cada blob desde el .git, lo convierte a texto (con \f entre páginas en los PDF)
y deja un manifiesto con páginas, caracteres y si parece escaneado.
"""
import json, os, subprocess, sys, tempfile, zipfile, re, html
from multiprocessing import Pool

REPO = '/home/user/libros-referencia3'
OUT = '/home/user/trabajo/corpus_ref3'
MANIF = '/home/user/trabajo/corpus_ref3/_manifiesto.jsonl'


def listar():
    out = subprocess.run(['git', '-C', REPO, 'ls-tree', '-r', '-z', '-l', 'HEAD'],
                         capture_output=True).stdout.decode('utf-8')
    items = []
    for rec in out.split('\0'):
        if not rec:
            continue
        meta, path = rec.split('\t', 1)
        _, _, sha, size = meta.split()
        items.append((path, sha, int(size)))
    return items


def pdf_a_texto(src, dst):
    subprocess.run(['pdftotext', '-enc', 'UTF-8', src, dst], capture_output=True, timeout=1800)
    pages = 0
    try:
        info = subprocess.run(['pdfinfo', src], capture_output=True, text=True, timeout=120).stdout
        m = re.search(r'^Pages:\s+(\d+)', info, re.M)
        pages = int(m.group(1)) if m else 0
    except Exception:
        pass
    return pages


def docx_a_texto(src, dst):
    import docx
    d = docx.Document(src)
    partes = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            partes.append(' | '.join(c.text for c in row.cells))
    open(dst, 'w').write('\n'.join(partes))


def pptx_a_texto(src, dst):
    import pptx
    pr = pptx.Presentation(src)
    partes = []
    for i, s in enumerate(pr.slides, 1):
        partes.append(f'\f[Diapositiva {i}]')
        for sh in s.shapes:
            if sh.has_text_frame:
                partes.append(sh.text_frame.text)
    open(dst, 'w').write('\n'.join(partes))


def epub_a_texto(src, dst):
    z = zipfile.ZipFile(src)
    partes = []
    for n in sorted(z.namelist()):
        if n.lower().endswith(('.xhtml', '.html', '.htm')):
            t = z.read(n).decode('utf-8', 'ignore')
            t = re.sub(r'<(script|style)[^>]*>.*?</\1>', ' ', t, flags=re.S)
            t = re.sub(r'<br\s*/?>|</p>|</h\d>|</div>', '\n', t)
            t = re.sub(r'<[^>]+>', '', t)
            partes.append('\f' + html.unescape(t))
    open(dst, 'w').write('\n'.join(partes))


def doc_a_texto(src, dst, tmpdir):
    subprocess.run(['soffice', '--headless', f'-env:UserInstallation=file://{tmpdir}/lo',
                    '--convert-to', 'txt:Text (encoded):UTF8', '--outdir', tmpdir, src],
                   capture_output=True, timeout=600)
    base = os.path.splitext(os.path.basename(src))[0] + '.txt'
    p = os.path.join(tmpdir, base)
    if os.path.exists(p):
        os.replace(p, dst)


def procesar(item):
    path, sha, size = item
    ext = os.path.splitext(path)[1].lower()
    dst = os.path.join(OUT, path + '.txt')
    os.makedirs(os.path.dirname(dst), exist_ok=True)
    rec = {'path': path, 'sha': sha, 'size': size, 'ext': ext, 'pages': 0, 'chars': 0, 'error': None}
    if os.path.exists(dst) and os.path.getsize(dst) > 0:
        rec['chars'] = len(open(dst, errors='ignore').read())
        rec['cached'] = True
    else:
        with tempfile.TemporaryDirectory(dir='/home/user/trabajo/tmp') as td:
            src = os.path.join(td, 'in' + ext)
            with open(src, 'wb') as f:
                subprocess.run(['git', '-C', REPO, 'cat-file', 'blob', sha], stdout=f, check=True)
            try:
                if ext == '.pdf':
                    rec['pages'] = pdf_a_texto(src, dst)
                elif ext == '.docx':
                    docx_a_texto(src, dst)
                elif ext == '.pptx':
                    pptx_a_texto(src, dst)
                elif ext == '.epub':
                    epub_a_texto(src, dst)
                elif ext == '.doc':
                    doc_a_texto(src, dst, td)
                elif ext in ('.md', '.txt'):
                    open(dst, 'wb').write(open(src, 'rb').read())
                else:
                    rec['error'] = 'tipo no soportado'
            except Exception as e:
                rec['error'] = repr(e)[:300]
        if os.path.exists(dst):
            rec['chars'] = len(open(dst, errors='ignore').read())
    if ext == '.pdf' and not rec['pages'] and os.path.exists(dst):
        rec['pages'] = open(dst, errors='ignore').read().count('\f')
    rec['chars_por_pagina'] = round(rec['chars'] / rec['pages']) if rec['pages'] else None
    rec['escaneado'] = bool(ext == '.pdf' and rec['pages'] and rec['chars'] / rec['pages'] < 300)
    return rec


if __name__ == '__main__':
    os.makedirs('/home/user/trabajo/tmp', exist_ok=True)
    items = listar()
    print(len(items), 'archivos', flush=True)
    with Pool(4) as pool, open(MANIF, 'w') as mf:
        for i, rec in enumerate(pool.imap_unordered(procesar, items, chunksize=1), 1):
            mf.write(json.dumps(rec, ensure_ascii=False) + '\n')
            mf.flush()
            if i % 25 == 0:
                print(i, 'hechos', flush=True)
    print('FIN', flush=True)
